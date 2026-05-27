"""DOCX report renderer mirroring the markdown structure. See spec §9.2.

This module turns an :class:`~byline.models.AuditResult` into a Microsoft Word
document with the same 11-section structure as the canonical Markdown report
(see :mod:`byline.report`). Framing is strictly comparative; the verbatim
disclaimer from spec §9.3 appears in the body, and a one-line disclaimer
fragment is set in every page header so it cannot be lost when an individual
page is printed or shared.
"""

from __future__ import annotations

import math
from collections.abc import Iterable
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

import byline
from byline.models import (
    AlignmentFindings,
    AuditResult,
    BaselineCorpus,
    BoilerplateFinding,
    ComparativeDelta,
    DisproportionFinding,
    FingerprintHit,
    HistoryFindings,
    SelfBaselineFinding,
    StyleProfile,
    VoiceFinding,
)

# ---------------------------------------------------------------------------
# Constants
# ---------------------------------------------------------------------------

#: Canonical disclaimer from spec §9.3 — quoted verbatim.
DISCLAIMER: str = (
    "This report presents stylistic signals comparing a candidate's submission "
    "to their own observable writing baseline. It is one input into a hiring "
    "decision, never a determination of authorship, and must not be treated as "
    "evidence of misconduct. False positives are possible — non-native English "
    "writers, proofread submissions, tutorial-derived code, and team-authored "
    "repos can all produce divergent signals."
)

#: One-sentence disclaimer used in the page header. Derived from the first
#: clause of the canonical disclaimer; kept short enough for a running header.
HEADER_DISCLAIMER: str = (
    "This report presents stylistic signals; it is one input into a hiring "
    "decision, never a determination of authorship."
)

#: StyleProfile fields, in the canonical emit order used by deltas and bullets.
_STYLE_PROFILE_FIELDS: tuple[str, ...] = (
    "em_dash_density",
    "emoji_in_headers_ratio",
    "avg_sentence_length",
    "type_token_ratio",
    "typo_rate",
    "sophistication_score",
    "banner_comment_density",
    "progress_ux_score",
)

#: Maximum length for a fingerprint excerpt before we truncate with an ellipsis.
_MAX_EXCERPT_LEN: int = 120

#: One-line plain-English gloss per overall_signal level.
_OVERALL_GLOSS: dict[str, str] = {
    "aligned": (
        "Target writing surface is broadly consistent with the candidate's observed baseline."
    ),
    "mixed": (
        "Some metrics diverge from baseline while others align; treat as a "
        "soft signal worth a closer look."
    ),
    "divergent": (
        "Multiple stylistic indicators diverge meaningfully from baseline; "
        "warrants a follow-up conversation."
    ),
    "highly_divergent": (
        "Stylistic indicators diverge sharply from baseline across multiple "
        "axes; warrants a thorough follow-up."
    ),
}


# ---------------------------------------------------------------------------
# Public API
# ---------------------------------------------------------------------------


def render_docx(result: AuditResult, output_path: Path) -> None:
    """Render a comparative-signal report as a Word ``.docx`` file.

    Produces the 11-section structure described in spec §9.2 with Arial as the
    document font, US Letter page size, a running header carrying the
    one-sentence disclaimer, a running footer with the byline version, and the
    full verbatim disclaimer from §9.3 inline near the top of the body.

    Args:
        result: The completed audit result to render.
        output_path: Filesystem path the ``.docx`` should be written to. The
            parent directory is created if it does not already exist.
    """
    output_path.parent.mkdir(parents=True, exist_ok=True)

    document = Document()
    _configure_document(document)
    _add_header_and_footer(document)

    _add_title_block(document, result)
    _add_disclaimer_callout(document)
    # v0.2: positive AI-disclosure callout, inserted before overall signal so
    # reviewers see the candidate's transparency first.
    if result.voice is not None and result.voice.ai_disclosure_found:
        _add_ai_disclosure_callout(document, result.voice)
    _add_overall_signal(document, result)
    _add_baseline_profile(document, result.baseline)
    _add_target_profile(document, result.target_profile)
    _add_deltas(document, result.deltas)
    _add_fingerprints(document, result.fingerprints)
    _add_disproportions(document, result.disproportions)
    # v0.2 findings — each section appears only when its field is not None.
    if result.history is not None:
        _add_history(document, result.history)
    if result.alignment is not None:
        _add_alignment(document, result.alignment)
    if result.voice is not None:
        _add_voice(document, result.voice)
    if result.boilerplate is not None:
        _add_boilerplate(document, result.boilerplate)
    if result.self_baseline is not None:
        _add_self_baseline(document, result.self_baseline)
    _add_qualitative(document, result.llm_qualitative)
    # v0.2: if alignment was run deterministic-only, surface the --with-llm
    # pointer immediately after the qualitative section.
    if result.alignment is not None and result.alignment.deterministic_only:
        _add_skipped_llm_note(document)
    _add_methodology(document)

    document.save(str(output_path))


# ---------------------------------------------------------------------------
# Document-level configuration
# ---------------------------------------------------------------------------


def _configure_document(document: Document) -> None:
    """Set Arial as the global font and US Letter as the page size."""
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(11)

    for section in document.sections:
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)


def _add_header_and_footer(document: Document) -> None:
    """Set the per-page header (disclaimer fragment) and footer (version line)."""
    for section in document.sections:
        header = section.header
        header_para = header.paragraphs[0]
        header_para.text = ""
        header_run = header_para.add_run(HEADER_DISCLAIMER)
        header_run.italic = True
        header_run.font.size = Pt(9)
        header_run.font.name = "Arial"

        footer = section.footer
        footer_para = footer.paragraphs[0]
        footer_para.text = ""
        footer_run = footer_para.add_run(
            f"byline v{byline.__version__} — comparative attribution analysis"
        )
        footer_run.italic = True
        footer_run.font.size = Pt(9)
        footer_run.font.name = "Arial"


# ---------------------------------------------------------------------------
# Section builders
# ---------------------------------------------------------------------------


def _add_title_block(document: Document, result: AuditResult) -> None:
    """Centered title plus a small subtitle of repo/candidate/timestamp."""
    title_para = document.add_paragraph()
    title_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_para.add_run("byline — Comparative Attribution Report")
    title_run.bold = True
    title_run.font.size = Pt(20)
    title_run.font.name = "Arial"

    candidate = result.candidate if result.candidate else "(not provided)"
    generated = result.generated_at.isoformat(timespec="seconds")

    subtitle_para = document.add_paragraph()
    subtitle_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle_text = (
        f"Target repo: {result.target_repo}    Candidate: {candidate}    Generated: {generated}"
    )
    subtitle_run = subtitle_para.add_run(subtitle_text)
    subtitle_run.font.size = Pt(10)
    subtitle_run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
    subtitle_run.font.name = "Arial"


def _add_disclaimer_callout(document: Document) -> None:
    """Inline disclaimer block — bold ``Disclaimer:`` prefix then verbatim text."""
    para = document.add_paragraph()
    prefix_run = para.add_run("Disclaimer: ")
    prefix_run.bold = True
    prefix_run.font.name = "Arial"

    body_run = para.add_run(DISCLAIMER)
    body_run.font.name = "Arial"


def _add_overall_signal(document: Document, result: AuditResult) -> None:
    document.add_heading("Overall signal", level=1)
    level = result.overall_signal
    label = level.replace("_", " ")
    gloss = _OVERALL_GLOSS.get(level, "")

    para = document.add_paragraph()
    head_run = para.add_run(f"Overall signal: {label}. ")
    head_run.bold = True
    head_run.font.name = "Arial"
    if gloss:
        body_run = para.add_run(gloss)
        body_run.font.name = "Arial"


def _add_baseline_profile(document: Document, baseline: BaselineCorpus | None) -> None:
    document.add_heading("Baseline profile", level=1)
    if baseline is None:
        para = document.add_paragraph(
            "No candidate baseline was supplied. Comparative deltas are "
            "unavailable; treat fingerprint and disproportion findings as "
            "standalone signals."
        )
        _apply_arial(para)
        return

    repos_count = len(baseline.repos_scanned)
    para = document.add_paragraph(
        f"Baseline aggregated from {repos_count} repo(s), {baseline.total_words} total words."
    )
    _apply_arial(para)


def _add_target_profile(document: Document, profile: StyleProfile) -> None:
    document.add_heading("Target profile", level=1)
    intro = document.add_paragraph(
        "Target style profile computed from the repository's Markdown and prose comments."
    )
    _apply_arial(intro)
    for field in _STYLE_PROFILE_FIELDS:
        value = getattr(profile, field)
        bullet = document.add_paragraph(style="List Bullet")
        name_run = bullet.add_run(f"{field}: ")
        name_run.bold = True
        name_run.font.name = "Arial"
        value_run = bullet.add_run(f"{value:.3f}")
        value_run.font.name = "Arial"


def _add_deltas(document: Document, deltas: list[ComparativeDelta]) -> None:
    document.add_heading("Comparative deltas", level=1)
    if not deltas:
        para = document.add_paragraph("No deltas (no baseline). Skipping comparative analysis.")
        _apply_arial(para)
        return

    by_metric = {d.metric: d for d in deltas}
    ordered: list[ComparativeDelta] = [
        by_metric[m] for m in _STYLE_PROFILE_FIELDS if m in by_metric
    ] + [d for d in deltas if d.metric not in _STYLE_PROFILE_FIELDS]

    columns = ("Metric", "Baseline", "Target", "Δ absolute", "Δ relative", "Severity")
    table = document.add_table(rows=1 + len(ordered), cols=len(columns))
    try:
        table.style = "Light Grid Accent 1"
    except KeyError:
        # Fall back to whatever default the template provides — borders are
        # nice-to-have, not load-bearing.
        pass

    header_cells = table.rows[0].cells
    for cell, label in zip(header_cells, columns, strict=False):
        cell.text = ""
        para = cell.paragraphs[0]
        run = para.add_run(label)
        run.bold = True
        run.font.name = "Arial"

    for row_idx, delta in enumerate(ordered, start=1):
        cells = table.rows[row_idx].cells
        values: Iterable[str] = (
            delta.metric,
            f"{delta.baseline_value:.3f}",
            f"{delta.target_value:.3f}",
            _format_signed(delta.absolute_delta),
            _format_relative(delta.relative_delta),
            delta.severity,
        )
        for cell, text in zip(cells, values, strict=False):
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(text)
            run.font.name = "Arial"


def _add_fingerprints(document: Document, fingerprints: list[FingerprintHit]) -> None:
    document.add_heading("Fingerprint findings", level=1)
    if not fingerprints:
        para = document.add_paragraph("No fingerprint hits in the target.")
        _apply_arial(para)
        return

    by_file: dict[str, list[FingerprintHit]] = {}
    for fp in fingerprints:
        by_file.setdefault(fp.file_path, []).append(fp)

    for file_path in sorted(by_file):
        document.add_heading(file_path, level=2)
        for fp in by_file[file_path]:
            excerpt = _truncate(fp.excerpt, _MAX_EXCERPT_LEN)
            bullet = document.add_paragraph(style="List Bullet")
            cat_run = bullet.add_run(f"{fp.category}")
            cat_run.bold = True
            cat_run.font.name = "Arial"
            tail_run = bullet.add_run(f" — {fp.pattern}: {excerpt}")
            tail_run.font.name = "Arial"


def _add_disproportions(document: Document, findings: list[DisproportionFinding]) -> None:
    document.add_heading("Disproportion findings", level=1)
    if not findings:
        para = document.add_paragraph("No structural disproportions detected.")
        _apply_arial(para)
        return
    for f in findings:
        bullet = document.add_paragraph(style="List Bullet")
        name_run = bullet.add_run(f"{f.name}")
        name_run.bold = True
        name_run.font.name = "Arial"
        body_run = bullet.add_run(
            f" ({f.severity}): observed {f.observed:.3f} vs threshold "
            f"{f.threshold:.3f}. {f.description}"
        )
        body_run.font.name = "Arial"


def _add_qualitative(document: Document, text: str | None) -> None:
    document.add_heading("Qualitative interpretation", level=1)
    if text is None:
        para = document.add_paragraph("No LLM qualitative pass was requested or available.")
        _apply_arial(para)
        return
    para = document.add_paragraph(text)
    _apply_arial(para)


def _add_ai_disclosure_callout(document: Document, voice: VoiceFinding) -> None:
    """Positive-framed callout for explicit AI-use disclosure (spec v0.2 §8).

    Rendered above the overall-signal heading so reviewers encounter the
    candidate's transparency before the divergence narrative.
    """
    document.add_heading("AI use disclosure", level=1)

    file_label = voice.ai_disclosure_file or "the candidate's documentation"
    excerpt = voice.ai_disclosure_excerpt or ""

    lead_para = document.add_paragraph()
    prefix_run = lead_para.add_run("Positive signal. ")
    prefix_run.bold = True
    prefix_run.font.name = "Arial"
    body_run = lead_para.add_run(
        "The candidate explicitly acknowledged the use of AI tools in "
        f"{file_label}:"
    )
    body_run.font.name = "Arial"

    excerpt_para = document.add_paragraph(excerpt, style="Intense Quote")
    _apply_arial(excerpt_para)

    closing_para = document.add_paragraph(
        "This is a transparency indicator and shifts the overall signal "
        'toward "aligned". Many of the divergence signals below may be '
        "explained by the candidate's stated tool use."
    )
    _apply_arial(closing_para)


def _add_history(document: Document, history: HistoryFindings) -> None:
    """Commit history forensics — timeline, messages, identity, file evolution."""
    document.add_heading("Commit history forensics", level=1)

    timeline = history.timeline
    messages = history.messages
    identity = history.identity

    span_days = timeline.span_seconds / 86400.0
    bursty_word = "bursty" if timeline.bursty else "not bursty"
    pasted_word = (
        "Appears pasted." if timeline.first_commit_appears_pasted else "Looks iterated."
    )
    drift_word = "Drift detected." if identity.drift_detected else "Consistent identity."

    para1 = document.add_paragraph(
        f"Timeline: {timeline.total_commits} commits over {span_days:.1f} days. "
        f"Burst density {timeline.burst_density:.2f}; {bursty_word}."
    )
    _apply_arial(para1)

    para2 = document.add_paragraph(
        f"First commit: {timeline.first_commit_file_count} files, "
        f"{timeline.first_commit_loc} LOC. {pasted_word}"
    )
    _apply_arial(para2)

    para3 = document.add_paragraph(
        f"Commit messages: {messages.total_messages} non-merge messages, "
        f"average length {messages.avg_length_chars:.0f} chars. "
        f"Debug-style commit ratio: {messages.debug_commit_ratio:.0%}. "
        f"Self-baseline divergence (vs. README): "
        f"{messages.self_baseline_divergence:.3f}."
    )
    _apply_arial(para3)

    para4 = document.add_paragraph(
        f"Author identity: {len(identity.unique_author_emails)} distinct "
        f"email(s). {drift_word}"
    )
    _apply_arial(para4)

    if history.file_evolutions:
        document.add_heading("File evolution", level=2)
        columns = ("File", "Commits", "Largest single addition", "Pasted?")
        table = document.add_table(
            rows=1 + len(history.file_evolutions), cols=len(columns)
        )
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:
            pass

        header_cells = table.rows[0].cells
        for cell, label in zip(header_cells, columns, strict=False):
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(label)
            run.bold = True
            run.font.name = "Arial"

        for row_idx, fe in enumerate(history.file_evolutions, start=1):
            cells = table.rows[row_idx].cells
            values: Iterable[str] = (
                fe.file_path,
                str(fe.total_commits_touching),
                str(fe.largest_single_addition_lines),
                "yes" if fe.appears_pasted else "no",
            )
            for cell, text in zip(cells, values, strict=False):
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(text)
                run.font.name = "Arial"


def _add_alignment(document: Document, alignment: AlignmentFindings) -> None:
    """Documentation-implementation alignment — mode, overall, checks table."""
    document.add_heading("Documentation-implementation alignment", level=1)
    mode = "deterministic-only" if alignment.deterministic_only else "deterministic + semantic"

    mode_para = document.add_paragraph(f"Mode: {mode}")
    _apply_arial(mode_para)

    overall_para = document.add_paragraph()
    label_run = overall_para.add_run("Overall: ")
    label_run.font.name = "Arial"
    value_run = overall_para.add_run(alignment.overall_alignment)
    value_run.bold = True
    value_run.font.name = "Arial"

    if alignment.llm_summary:
        summary_para = document.add_paragraph(alignment.llm_summary)
        _apply_arial(summary_para)

    if alignment.checks:
        document.add_heading("Checks", level=2)
        columns = ("Kind", "Source", "Severity", "Description")
        table = document.add_table(rows=1 + len(alignment.checks), cols=len(columns))
        try:
            table.style = "Light Grid Accent 1"
        except KeyError:
            pass

        header_cells = table.rows[0].cells
        for cell, label in zip(header_cells, columns, strict=False):
            cell.text = ""
            para = cell.paragraphs[0]
            run = para.add_run(label)
            run.bold = True
            run.font.name = "Arial"

        for row_idx, check in enumerate(alignment.checks, start=1):
            cells = table.rows[row_idx].cells
            values: Iterable[str] = (
                check.kind,
                check.source,
                check.severity,
                check.description,
            )
            for cell, text in zip(cells, values, strict=False):
                cell.text = ""
                para = cell.paragraphs[0]
                run = para.add_run(text)
                run.font.name = "Arial"


def _add_voice(document: Document, voice: VoiceFinding) -> None:
    """Voice-and-disclosure section — first-person voice + AI-use disclosure."""
    document.add_heading("Voice and disclosure", level=1)

    presence = "Present." if voice.has_first_person_voice else "Sparse."
    para1 = document.add_paragraph(
        f"First-person voice: {voice.first_person_count} occurrences "
        f"({voice.first_person_per_1k_words:.1f} per 1000 words). {presence}"
    )
    _apply_arial(para1)

    disclosure = (
        "Explicit AI-use disclosure found."
        if voice.ai_disclosure_found
        else "No explicit AI-use disclosure detected."
    )
    para2 = document.add_paragraph(disclosure)
    _apply_arial(para2)


def _add_boilerplate(document: Document, boilerplate: BoilerplateFinding) -> None:
    """Boilerplate-density section — meta-files present out of those checked."""
    document.add_heading("Boilerplate density", level=1)

    n_present = len(boilerplate.meta_files_present)
    n_checked = len(boilerplate.meta_files_checked)
    summary_para = document.add_paragraph()
    text_run = summary_para.add_run(
        f"{n_present} of {n_checked} standard meta-files present "
        f"({boilerplate.density_ratio:.0%}). Severity: "
    )
    text_run.font.name = "Arial"
    sev_run = summary_para.add_run(boilerplate.severity)
    sev_run.bold = True
    sev_run.font.name = "Arial"
    tail_run = summary_para.add_run(".")
    tail_run.font.name = "Arial"

    present_list = ", ".join(boilerplate.meta_files_present) or "none"
    present_para = document.add_paragraph(f"Present: {present_list}")
    _apply_arial(present_para)


def _add_self_baseline(document: Document, self_baseline: SelfBaselineFinding) -> None:
    """Self-baseline within-repo divergence — commit msg vs README, code vs README."""
    document.add_heading("Self-baseline within-repo divergence", level=1)

    msg_para = document.add_paragraph(
        f"Commit messages vs. README: {self_baseline.commit_msg_vs_readme_distance:.3f}"
    )
    _apply_arial(msg_para)
    code_para = document.add_paragraph(
        f"Code comments vs. README: {self_baseline.code_comment_vs_readme_distance:.3f}"
    )
    _apply_arial(code_para)

    overall_para = document.add_paragraph()
    label_run = overall_para.add_run("Within-repo divergence: ")
    label_run.font.name = "Arial"
    value_run = overall_para.add_run(self_baseline.within_repo_divergence)
    value_run.bold = True
    value_run.font.name = "Arial"

    note_para = document.add_paragraph(self_baseline.note)
    _apply_arial(note_para)


def _add_skipped_llm_note(document: Document) -> None:
    """One-line pointer when alignment was deterministic-only."""
    para = document.add_paragraph(
        "Semantic alignment and qualitative summary skipped. "
        "Run with --with-llm to enable.",
        style="Intense Quote",
    )
    _apply_arial(para)


def _add_methodology(document: Document) -> None:
    document.add_heading("Methodology", level=1)
    para = document.add_paragraph(
        "Signals are computed by comparing eight stylistic metrics on the "
        "target repository against the candidate's aggregated writing "
        "baseline, then cross-referenced against a catalogue of known "
        "phrasing and structural patterns. Severities follow fixed "
        "thresholds; nothing here is a verdict."
    )
    _apply_arial(para)

    link_para = document.add_paragraph(
        "See docs/methodology.md for full metric definitions, thresholds, and limitations."
    )
    _apply_arial(link_para)


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------


def _apply_arial(paragraph) -> None:  # type: ignore[no-untyped-def]
    """Force every run in a paragraph to the Arial face."""
    for run in paragraph.runs:
        run.font.name = "Arial"


def _format_signed(value: float) -> str:
    """Format a float with an explicit sign and 3 decimals."""
    sign = "+" if value >= 0 else ""
    return f"{sign}{value:.3f}"


def _format_relative(value: float) -> str:
    """Format a relative delta as a percentage, with sentinels for extremes."""
    if math.isinf(value):
        return "+inf" if value > 0 else "-inf"
    if math.isnan(value):
        return "n/a"
    if abs(value) < 5:
        sign = "+" if value >= 0 else ""
        return f"{sign}{value * 100:.1f}%"
    sign = "+" if value >= 0 else "-"
    return f"{sign}>500%"


def _truncate(text: str, limit: int) -> str:
    """Truncate ``text`` to at most ``limit`` chars, appending an ellipsis."""
    if len(text) <= limit:
        return text
    return text[: limit - 1].rstrip() + "…"


__all__ = ["DISCLAIMER", "HEADER_DISCLAIMER", "render_docx"]
