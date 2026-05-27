"""Tests for byline.report — markdown comparative-signal report rendering per spec §9.1/§9.3."""

from __future__ import annotations

from datetime import datetime, timezone

import byline
from byline.models import (
    AuditResult,
    BaselineCorpus,
    ComparativeDelta,
    DisproportionFinding,
    FingerprintHit,
    StyleProfile,
    WritingSample,
)
from byline.report import render_markdown

# ---------------------------------------------------------------------------
# Canonical disclaimer text from spec §9.3 — must appear verbatim in output.
# ---------------------------------------------------------------------------

DISCLAIMER = (
    "This report presents stylistic signals comparing a candidate's submission "
    "to their own observable writing baseline. It is one input into a hiring "
    "decision, never a determination of authorship, and must not be treated as "
    "evidence of misconduct. False positives are possible — non-native English "
    "writers, proofread submissions, tutorial-derived code, and team-authored "
    "repos can all produce divergent signals."
)


# ---------------------------------------------------------------------------
# Fixture builders
# ---------------------------------------------------------------------------


def _zero_profile() -> StyleProfile:
    return StyleProfile(
        em_dash_density=0.0,
        emoji_in_headers_ratio=0.0,
        avg_sentence_length=0.0,
        type_token_ratio=0.0,
        typo_rate=0.0,
        sophistication_score=0.0,
        banner_comment_density=0.0,
        progress_ux_score=0.0,
    )


def _profile(**overrides: float) -> StyleProfile:
    base = dict(
        em_dash_density=2.0,
        emoji_in_headers_ratio=0.1,
        avg_sentence_length=15.0,
        type_token_ratio=0.5,
        typo_rate=1.0,
        sophistication_score=0.3,
        banner_comment_density=1.0,
        progress_ux_score=0.25,
    )
    base.update(overrides)
    return StyleProfile(**base)


def _baseline() -> BaselineCorpus:
    return BaselineCorpus(
        username="alice",
        samples=[
            WritingSample(
                source="README.md@alice/proj-one",
                kind="readme",
                text="hello world",
                word_count=2,
            ),
            WritingSample(
                source="README.md@alice/proj-two",
                kind="readme",
                text="another sample",
                word_count=2,
            ),
        ],
        total_words=4,
        repos_scanned=["alice/proj-one", "alice/proj-two"],
        repos_skipped=[],
    )


def _deltas() -> list[ComparativeDelta]:
    return [
        ComparativeDelta(
            metric="em_dash_density",
            baseline_value=0.0,
            target_value=12.5,
            absolute_delta=12.5,
            relative_delta=float("inf"),
            severity="extreme",
        ),
        ComparativeDelta(
            metric="emoji_in_headers_ratio",
            baseline_value=0.1,
            target_value=0.6,
            absolute_delta=0.5,
            relative_delta=5.0,
            severity="significant",
        ),
        ComparativeDelta(
            metric="avg_sentence_length",
            baseline_value=15.0,
            target_value=16.0,
            absolute_delta=1.0,
            relative_delta=0.0667,
            severity="aligned",
        ),
    ]


def _fingerprints() -> list[FingerprintHit]:
    return [
        FingerprintHit(
            file_path="README.md",
            line_number=12,
            pattern="dive into",
            category="phrase",
            excerpt="Let's dive into this project",
        ),
        FingerprintHit(
            file_path="README.md",
            line_number=18,
            pattern="elevate",
            category="phrase",
            excerpt="This will elevate your workflow to new heights",
        ),
        FingerprintHit(
            file_path="setup.sh",
            line_number=1,
            pattern="banner-comment-block",
            category="shell_banner",
            excerpt="#" * 40,
        ),
    ]


def _disproportions() -> list[DisproportionFinding]:
    return [
        DisproportionFinding(
            name="docs_to_code_ratio",
            observed=4.2,
            threshold=2.0,
            severity="significant",
            description="Documentation lines exceed code lines by an unusual margin.",
        ),
        DisproportionFinding(
            name="comment_to_code_ratio",
            observed=0.85,
            threshold=0.5,
            severity="notable",
            description="Comment density is higher than typical for the language.",
        ),
    ]


def _make_result(
    *,
    baseline: BaselineCorpus | None = None,
    deltas: list[ComparativeDelta] | None = None,
    fingerprints: list[FingerprintHit] | None = None,
    disproportions: list[DisproportionFinding] | None = None,
    llm_qualitative: str | None = None,
    overall: str = "divergent",
) -> AuditResult:
    return AuditResult(
        target_repo="alice/take-home-submission",
        candidate="alice",
        baseline=baseline,
        target_profile=_profile(em_dash_density=12.5, emoji_in_headers_ratio=0.6),
        fingerprints=fingerprints if fingerprints is not None else _fingerprints(),
        disproportions=disproportions if disproportions is not None else _disproportions(),
        deltas=deltas if deltas is not None else _deltas(),
        llm_qualitative=llm_qualitative,
        overall_signal=overall,  # type: ignore[arg-type]
        generated_at=datetime(2026, 5, 27, 12, 0, 0, tzinfo=timezone.utc),
    )


# ---------------------------------------------------------------------------
# Required-sections / structural tests
# ---------------------------------------------------------------------------


REQUIRED_HEADERS = [
    "# byline — Comparative Attribution Report",
    "## Overall signal",
    "## Baseline profile",
    "## Target profile",
    "## Comparative deltas",
    "## Fingerprint findings",
    "## Disproportion findings",
    "## Qualitative interpretation",
    "## Methodology",
]


def test_render_markdown_includes_all_required_section_headers() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    for header in REQUIRED_HEADERS:
        assert header in md, f"missing header: {header!r}"


def test_render_markdown_sections_appear_in_spec_order() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    positions = [md.index(h) for h in REQUIRED_HEADERS]
    assert positions == sorted(positions), f"sections out of order: {positions}"


def test_render_markdown_includes_verbatim_disclaimer() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert DISCLAIMER in md


def test_render_markdown_includes_target_repo_and_candidate() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert "alice/take-home-submission" in md
    assert "alice" in md


def test_render_markdown_includes_byline_version() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert byline.__version__ in md


def test_render_markdown_has_no_banned_phrases() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result).lower()
    assert "ai detector" not in md
    assert "detect ai" not in md


# ---------------------------------------------------------------------------
# Empty-state handling
# ---------------------------------------------------------------------------


def test_render_markdown_no_baseline_says_baseline_missing() -> None:
    result = _make_result(baseline=None, deltas=[])
    md = render_markdown(result)
    assert "No candidate baseline was supplied" in md


def test_render_markdown_no_deltas_says_skipping() -> None:
    result = _make_result(baseline=None, deltas=[])
    md = render_markdown(result)
    assert "No deltas" in md


def test_render_markdown_no_fingerprints_handled() -> None:
    result = _make_result(baseline=_baseline(), fingerprints=[])
    md = render_markdown(result)
    # Section header is still emitted; an empty-state phrase follows.
    assert "## Fingerprint findings" in md
    # Should not crash, should not render any "###" file subheader.
    # (find a stable empty-state token — implementation will say "No fingerprint")
    assert "No fingerprint" in md or "no fingerprint" in md.lower()


def test_render_markdown_no_disproportions_handled() -> None:
    result = _make_result(baseline=_baseline(), disproportions=[])
    md = render_markdown(result)
    assert "## Disproportion findings" in md


def test_render_markdown_with_llm_qualitative_includes_text() -> None:
    qual = "The target prose reads as substantially more formal than the baseline."
    result = _make_result(baseline=_baseline(), llm_qualitative=qual)
    md = render_markdown(result)
    assert qual in md


def test_render_markdown_without_llm_qualitative_says_unavailable() -> None:
    result = _make_result(baseline=_baseline(), llm_qualitative=None)
    md = render_markdown(result)
    assert "No LLM qualitative pass" in md


# ---------------------------------------------------------------------------
# Content / formatting details
# ---------------------------------------------------------------------------


def test_render_markdown_deltas_table_has_header_row() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert "| Metric | Baseline | Target | Δ absolute | Δ relative | Severity |" in md


def test_render_markdown_fingerprints_grouped_by_file_path() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert "### README.md" in md
    assert "### setup.sh" in md


def test_render_markdown_disclaimer_in_blockquote() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    # Blockquoted form of the first sentence-fragment must exist.
    assert "> This report presents stylistic signals" in md


def test_render_markdown_methodology_link_present() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert "docs/methodology.md" in md


def test_render_markdown_includes_generated_at_timestamp() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    # ISO-formatted with second precision.
    assert "2026-05-27T12:00:00" in md


def test_render_markdown_returns_str() -> None:
    result = _make_result(baseline=_baseline())
    md = render_markdown(result)
    assert isinstance(md, str)
    assert len(md) > 0


def test_render_markdown_handles_long_fingerprint_excerpt_truncation() -> None:
    long_excerpt = "x" * 500
    fps = [
        FingerprintHit(
            file_path="big.md",
            line_number=1,
            pattern="long",
            category="phrase",
            excerpt=long_excerpt,
        )
    ]
    result = _make_result(baseline=_baseline(), fingerprints=fps)
    md = render_markdown(result)
    # Full 500-char string must not appear; truncation marker should.
    assert long_excerpt not in md
    assert "…" in md


# ---------------------------------------------------------------------------
# DOCX rendering — spec §9.2
# ---------------------------------------------------------------------------


from byline.report_docx import HEADER_DISCLAIMER, render_docx  # noqa: E402


def _docx_full_body_text(doc) -> str:  # type: ignore[no-untyped-def]
    """Return all paragraph + table-cell text from a docx Document as one blob."""
    parts: list[str] = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def _docx_header_footer_text(doc) -> str:  # type: ignore[no-untyped-def]
    parts: list[str] = []
    for section in doc.sections:
        for para in section.header.paragraphs:
            parts.append(para.text)
        for para in section.footer.paragraphs:
            parts.append(para.text)
    return "\n".join(parts)


def test_render_docx_writes_non_empty_file(tmp_path) -> None:
    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)
    assert out.exists()
    assert out.stat().st_size > 1000


def test_render_docx_disclaimer_verbatim_in_body(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert DISCLAIMER in body


def test_render_docx_no_banned_phrases(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    full = (_docx_full_body_text(doc) + "\n" + _docx_header_footer_text(doc)).lower()
    assert "ai detector" not in full
    assert "detect ai" not in full


def test_render_docx_table_present_when_deltas_exist(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    assert len(doc.tables) >= 1, "expected at least one deltas table"

    table = doc.tables[0]
    header_row_text = " | ".join(cell.text for cell in table.rows[0].cells)
    assert "Metric" in header_row_text
    assert "Baseline" in header_row_text
    assert "Target" in header_row_text
    assert "Severity" in header_row_text
    # 1 header row + one row per delta from _deltas()
    assert len(table.rows) == 1 + len(_deltas())


def test_render_docx_no_baseline_no_table_but_explanation(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=None, deltas=[])
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert "No candidate baseline was supplied" in body
    assert "No deltas" in body
    # When there are no deltas there should be no deltas table.
    assert len(doc.tables) == 0


def test_render_docx_includes_all_required_section_headers(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    for heading in (
        "Overall signal",
        "Baseline profile",
        "Target profile",
        "Comparative deltas",
        "Fingerprint findings",
        "Disproportion findings",
        "Qualitative interpretation",
        "Methodology",
    ):
        assert heading in body, f"missing section heading: {heading!r}"


def test_render_docx_header_and_footer_set(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    hf_text = _docx_header_footer_text(doc)
    assert HEADER_DISCLAIMER in hf_text
    assert f"byline v{byline.__version__}" in hf_text


def test_render_docx_normal_style_uses_arial(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    assert doc.styles["Normal"].font.name == "Arial"


def test_render_docx_us_letter_page_size(tmp_path) -> None:
    from docx import Document
    from docx.shared import Inches

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    section = doc.sections[0]
    assert section.page_width == Inches(8.5)
    assert section.page_height == Inches(11)


def test_render_docx_title_present(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert "byline — Comparative Attribution Report" in body
    # Subtitle metadata appears too.
    assert "alice/take-home-submission" in body
    assert "2026-05-27T12:00:00" in body


def test_render_docx_creates_parent_directories(tmp_path) -> None:
    result = _make_result(baseline=_baseline())
    nested = tmp_path / "nested" / "deeper" / "report.docx"
    assert not nested.parent.exists()
    render_docx(result, nested)
    assert nested.exists()


def test_render_docx_methodology_text_present(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline())
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert "docs/methodology.md" in body


def test_render_docx_with_llm_qualitative_includes_text(tmp_path) -> None:
    from docx import Document

    qual = "Prose reads more formal than the baseline; vocabulary mix narrower."
    result = _make_result(baseline=_baseline(), llm_qualitative=qual)
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert qual in body


def test_render_docx_without_llm_qualitative_says_unavailable(tmp_path) -> None:
    from docx import Document

    result = _make_result(baseline=_baseline(), llm_qualitative=None)
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    assert "No LLM qualitative pass" in body


# ---------------------------------------------------------------------------
# v0.2 — new sections (history, alignment, voice, boilerplate, self_baseline)
# ---------------------------------------------------------------------------


from byline.models import (  # noqa: E402
    AlignmentCheck,
    AlignmentFindings,
    AuthorIdentityFinding,
    BoilerplateFinding,
    CommitMessageStyleFinding,
    CommitTimelineFinding,
    FileEvolutionFinding,
    HistoryFindings,
    SelfBaselineFinding,
    VoiceFinding,
)


def _history_findings() -> HistoryFindings:
    return HistoryFindings(
        timeline=CommitTimelineFinding(
            total_commits=20,
            span_seconds=86400 * 3,  # 3 days
            burst_density=0.85,
            burst_window_start=datetime(2026, 5, 25, 14, 0, 0, tzinfo=timezone.utc),
            bursty=True,
            first_commit_file_count=42,
            first_commit_loc=3500,
            first_commit_appears_pasted=True,
        ),
        messages=CommitMessageStyleFinding(
            total_messages=18,
            avg_length_chars=72.5,
            style_profile=_zero_profile(),
            debug_commit_ratio=0.05,
            self_baseline_divergence=0.42,
        ),
        identity=AuthorIdentityFinding(
            unique_author_emails=["alice@example.com", "ai-bot@example.com"],
            unique_author_names=["Alice", "Coding Assistant"],
            drift_detected=True,
        ),
        file_evolutions=[
            FileEvolutionFinding(
                file_path="src/main.py",
                total_commits_touching=4,
                largest_single_addition_lines=900,
                appears_pasted=True,
            ),
            FileEvolutionFinding(
                file_path="README.md",
                total_commits_touching=6,
                largest_single_addition_lines=80,
                appears_pasted=False,
            ),
        ],
    )


def _alignment_findings(deterministic_only: bool = True) -> AlignmentFindings:
    return AlignmentFindings(
        checks=[
            AlignmentCheck(
                kind="cli_flag_documented_missing_in_code",
                source="deterministic",
                description="README mentions --foo flag, not found in CLI.",
                doc_location="README.md:42",
                code_location=None,
                severity="notable",
            ),
            AlignmentCheck(
                kind="env_var_documented_missing_in_code",
                source="deterministic",
                description="DOCS_VAR documented but unused in source.",
                doc_location="README.md:88",
                code_location=None,
                severity="info",
            ),
        ],
        deterministic_only=deterministic_only,
        overall_alignment="minor_gaps",
        llm_summary=(
            None
            if deterministic_only
            else "Documentation broadly matches the implementation with minor gaps."
        ),
    )


def _voice_finding(ai_disclosure: bool = False) -> VoiceFinding:
    return VoiceFinding(
        first_person_count=12,
        first_person_per_1k_words=3.4,
        has_first_person_voice=True,
        ai_disclosure_found=ai_disclosure,
        ai_disclosure_file="AI_USE.md" if ai_disclosure else None,
        ai_disclosure_excerpt=(
            "I used Claude to draft the initial scaffolding and refined by hand."
            if ai_disclosure
            else None
        ),
    )


def _boilerplate_finding() -> BoilerplateFinding:
    return BoilerplateFinding(
        meta_files_present=["README.md", "LICENSE"],
        meta_files_checked=["README.md", "LICENSE", "CONTRIBUTING.md", ".gitignore"],
        density_ratio=0.5,
        severity="notable",
    )


def _self_baseline_finding() -> SelfBaselineFinding:
    return SelfBaselineFinding(
        commit_msg_vs_readme_distance=0.71,
        code_comment_vs_readme_distance=0.55,
        within_repo_divergence="notable",
        note="Commit messages read more terse than the README's expansive tone.",
    )


def _make_v02_result(
    *,
    history: HistoryFindings | None = None,
    alignment: AlignmentFindings | None = None,
    voice: VoiceFinding | None = None,
    boilerplate: BoilerplateFinding | None = None,
    self_baseline: SelfBaselineFinding | None = None,
    llm_qualitative: str | None = None,
) -> AuditResult:
    return AuditResult(
        target_repo="alice/take-home-submission",
        candidate="alice",
        baseline=_baseline(),
        target_profile=_profile(),
        fingerprints=_fingerprints(),
        disproportions=_disproportions(),
        deltas=_deltas(),
        llm_qualitative=llm_qualitative,
        overall_signal="divergent",
        generated_at=datetime(2026, 5, 27, 12, 0, 0, tzinfo=timezone.utc),
        history=history,
        alignment=alignment,
        voice=voice,
        boilerplate=boilerplate,
        self_baseline=self_baseline,
    )


# Markdown: AI-disclosure callout (positive framing)


def test_render_markdown_ai_disclosure_positive_callout() -> None:
    result = _make_v02_result(voice=_voice_finding(ai_disclosure=True))
    md = render_markdown(result)
    assert "## AI use disclosure" in md
    assert "Positive signal" in md
    assert "AI_USE.md" in md
    assert "I used Claude to draft" in md
    # Banned phrases must not appear in the AI-disclosure section.
    assert "ai detector" not in md.lower()
    assert "detect ai" not in md.lower()


def test_render_markdown_ai_disclosure_callout_appears_before_overall_signal() -> None:
    result = _make_v02_result(voice=_voice_finding(ai_disclosure=True))
    md = render_markdown(result)
    callout_pos = md.index("## AI use disclosure")
    overall_pos = md.index("## Overall signal")
    assert callout_pos < overall_pos


def test_render_markdown_no_ai_disclosure_no_callout() -> None:
    result = _make_v02_result(voice=_voice_finding(ai_disclosure=False))
    md = render_markdown(result)
    assert "## AI use disclosure" not in md


# Backwards compatibility — v0.1 audits (all v0.2 fields None) must produce
# identical section structure to before.


def test_render_markdown_v01_compatible_omits_new_sections() -> None:
    result = _make_v02_result()
    md = render_markdown(result)
    assert "## AI use disclosure" not in md
    assert "## Commit history forensics" not in md
    assert "## Documentation-implementation alignment" not in md
    assert "## Voice and disclosure" not in md
    assert "## Boilerplate density" not in md
    assert "## Self-baseline within-repo divergence" not in md


# History section


def test_render_markdown_history_section_present_when_populated() -> None:
    result = _make_v02_result(history=_history_findings())
    md = render_markdown(result)
    assert "## Commit history forensics" in md
    # Timeline summary content.
    assert "20 commits" in md
    assert "bursty" in md
    # First-commit summary mentions pasted-style observation.
    assert "Appears pasted" in md
    # Identity drift.
    assert "Drift detected" in md
    # File evolution table header.
    assert "| File | Commits | Largest single addition | Pasted? |" in md
    assert "src/main.py" in md
    assert "README.md" in md


def test_render_markdown_history_section_after_disproportions_before_qualitative() -> None:
    result = _make_v02_result(history=_history_findings())
    md = render_markdown(result)
    disp_pos = md.index("## Disproportion findings")
    hist_pos = md.index("## Commit history forensics")
    qual_pos = md.index("## Qualitative interpretation")
    assert disp_pos < hist_pos < qual_pos


# Alignment section


def test_render_markdown_alignment_section_deterministic_only() -> None:
    result = _make_v02_result(alignment=_alignment_findings(deterministic_only=True))
    md = render_markdown(result)
    assert "## Documentation-implementation alignment" in md
    assert "deterministic-only" in md
    assert "minor_gaps" in md
    assert "| Kind | Source | Severity | Description |" in md
    assert "cli_flag_documented_missing_in_code" in md


def test_render_markdown_alignment_section_with_semantic_summary() -> None:
    result = _make_v02_result(alignment=_alignment_findings(deterministic_only=False))
    md = render_markdown(result)
    assert "deterministic + semantic" in md
    assert "Documentation broadly matches the implementation" in md


# Voice section


def test_render_markdown_voice_section_present() -> None:
    result = _make_v02_result(voice=_voice_finding(ai_disclosure=False))
    md = render_markdown(result)
    assert "## Voice and disclosure" in md
    assert "12 occurrences" in md
    assert "3.4 per 1000 words" in md
    assert "No explicit AI-use disclosure detected" in md


# Boilerplate section


def test_render_markdown_boilerplate_section_present() -> None:
    result = _make_v02_result(boilerplate=_boilerplate_finding())
    md = render_markdown(result)
    assert "## Boilerplate density" in md
    assert "2 of 4 standard meta-files present" in md
    assert "50%" in md
    assert "README.md, LICENSE" in md


# Self-baseline section


def test_render_markdown_self_baseline_section_present() -> None:
    result = _make_v02_result(self_baseline=_self_baseline_finding())
    md = render_markdown(result)
    assert "## Self-baseline within-repo divergence" in md
    assert "0.710" in md
    assert "0.550" in md
    assert "notable" in md
    assert "Commit messages read more terse" in md


# Skipped-LLM note when alignment was run deterministic-only


def test_render_markdown_skipped_llm_note_when_alignment_deterministic_only() -> None:
    result = _make_v02_result(alignment=_alignment_findings(deterministic_only=True))
    md = render_markdown(result)
    assert "Run with `--with-llm` to enable" in md


def test_render_markdown_no_skipped_llm_note_when_alignment_with_semantic() -> None:
    result = _make_v02_result(alignment=_alignment_findings(deterministic_only=False))
    md = render_markdown(result)
    assert "Run with `--with-llm` to enable" not in md


# Full ordering of new sections


def test_render_markdown_v02_section_order_matches_spec() -> None:
    result = _make_v02_result(
        history=_history_findings(),
        alignment=_alignment_findings(deterministic_only=True),
        voice=_voice_finding(ai_disclosure=True),
        boilerplate=_boilerplate_finding(),
        self_baseline=_self_baseline_finding(),
    )
    md = render_markdown(result)
    expected_order = [
        "## AI use disclosure",
        "## Overall signal",
        "## Baseline profile",
        "## Target profile",
        "## Comparative deltas",
        "## Fingerprint findings",
        "## Disproportion findings",
        "## Commit history forensics",
        "## Documentation-implementation alignment",
        "## Voice and disclosure",
        "## Boilerplate density",
        "## Self-baseline within-repo divergence",
        "## Qualitative interpretation",
        "## Methodology",
    ]
    positions = [md.index(h) for h in expected_order]
    assert positions == sorted(positions), (
        f"sections out of order: expected {expected_order}, got positions {positions}"
    )


# DOCX rendering for v0.2


def test_render_docx_v02_full_sections(tmp_path) -> None:
    from docx import Document

    result = _make_v02_result(
        history=_history_findings(),
        alignment=_alignment_findings(deterministic_only=False),
        voice=_voice_finding(ai_disclosure=True),
        boilerplate=_boilerplate_finding(),
        self_baseline=_self_baseline_finding(),
    )
    out = tmp_path / "report.docx"
    render_docx(result, out)
    assert out.exists()
    assert out.stat().st_size > 1000

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    for heading in (
        "AI use disclosure",
        "Commit history forensics",
        "Documentation-implementation alignment",
        "Voice and disclosure",
        "Boilerplate density",
        "Self-baseline within-repo divergence",
    ):
        assert heading in body, f"missing v0.2 heading: {heading!r}"
    # Positive framing token must appear.
    assert "Positive signal" in body
    assert "AI_USE.md" in body
    # File evolution table content.
    assert "src/main.py" in body
    # Alignment summary text.
    assert "Documentation broadly matches the implementation" in body
    # Boilerplate present-list.
    assert "README.md" in body and "LICENSE" in body
    # Self-baseline note.
    assert "Commit messages read more terse" in body


def test_render_docx_v02_omits_sections_when_fields_none(tmp_path) -> None:
    from docx import Document

    result = _make_v02_result()  # all v0.2 fields None
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    body = _docx_full_body_text(doc)
    for missing in (
        "AI use disclosure",
        "Commit history forensics",
        "Documentation-implementation alignment",
        "Voice and disclosure",
        "Boilerplate density",
        "Self-baseline within-repo divergence",
    ):
        assert missing not in body, f"unexpected v0.2 heading present: {missing!r}"


def test_render_docx_v02_no_banned_phrases(tmp_path) -> None:
    from docx import Document

    result = _make_v02_result(
        history=_history_findings(),
        alignment=_alignment_findings(deterministic_only=True),
        voice=_voice_finding(ai_disclosure=True),
        boilerplate=_boilerplate_finding(),
        self_baseline=_self_baseline_finding(),
    )
    out = tmp_path / "report.docx"
    render_docx(result, out)

    doc = Document(str(out))
    full = (_docx_full_body_text(doc) + "\n" + _docx_header_footer_text(doc)).lower()
    assert "ai detector" not in full
    assert "detect ai" not in full
